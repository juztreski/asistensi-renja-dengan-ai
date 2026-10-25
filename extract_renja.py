"""Ekstraksi dokumen Renja/Rencana Kinerja OPD.

Mendukung PDF, DOCX, XLSX, dan CSV. Tabel Renja dipetakan secara heuristik ke
hierarki Program -> Kegiatan -> Subkegiatan. Bila tabel tidak terbaca (umumnya
PDF hasil pindai), teks akan diparse baris demi baris dan dapat dilanjutkan
dengan fallback LLM.
"""

from __future__ import annotations

import csv
import io
import os
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

CODE_RE = re.compile(r"(?<!\d)(\d{1,2}(?:\.\d{1,2}){1,4})(?!\d)")
SPACE_RE = re.compile(r"\s+")
YEAR_RE = re.compile(r"\b(20\d{2})\b")

LEVEL_PROGRAM = "Program"
LEVEL_KEGIATAN = "Kegiatan"
LEVEL_SUBKEGIATAN = "Subkegiatan"

HEADER_PATTERNS: dict[str, re.Pattern] = {
    "kode": re.compile(r"\b(kode|kd)\b", re.I),
    "nama": re.compile(r"(uraian|nomenklatur|nama)", re.I),
    "indikator": re.compile(r"indikator", re.I),
    "target": re.compile(r"\btarget\b", re.I),
    "satuan": re.compile(r"\bsatuan\b", re.I),
    "pagu": re.compile(r"(pagu|anggaran|biaya|jumlah)", re.I),
    "lokasi": re.compile(r"(lokasi|wilayah)", re.I),
    "keterangan": re.compile(r"(keterangan|\bket\b)", re.I),
    "level": re.compile(r"\b(level|tingkat|jenis)\b", re.I),
    "sasaran": re.compile(r"sasaran", re.I),
    "no": re.compile(r"^(no|nomor)\b", re.I),
}

SUMMARY_RE = re.compile(r"^(jumlah|total|sub\s*jumlah|grand\s*total)\b", re.I)
HEADERISH_RE = re.compile(
    r"(uraian|indikator|target|satuan|pagu|anggaran|kode|nomenklatur)", re.I
)
LEVEL_LABEL_RE = re.compile(r"^(program|kegiatan|sub\s*kegiatan|subkegiatan)\b[\s:\-]*", re.I)
LEADING_NUMBER_RE = re.compile(r"^\s*(?:[IVXLC]+\.|[A-Za-z]\.|\d+[\.\)])\s+")
PAGU_RE = re.compile(r"[^0-9,.\-]")

STOPWORDS = {
    "DAN",
    "ATAU",
    "YANG",
    "UNTUK",
    "DENGAN",
    "PADA",
    "DI",
    "KE",
    "DARI",
    "DALAM",
    "SEBAGAI",
    "SECARA",
    "SERTA",
    "JUMLAH",
    "PERSEN",
}


# ----------------------------------------------------------------------
# Utilitas umum
# ----------------------------------------------------------------------
def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def norm_space(value: Any) -> str:
    if value is None:
        return ""
    return SPACE_RE.sub(" ", str(value)).strip()


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).replace("\r", "\n").replace("\xa0", " ")
    return norm_space(text)


def extract_code(value: Any) -> str:
    text = clean_cell(value)
    if not text:
        return ""
    match = CODE_RE.search(text)
    return match.group(1) if match else ""


def strip_level_label(value: Any) -> tuple[str, str]:
    """Pisahkan label level di awal teks dan kembalikan (level, sisa teks)."""
    text = clean_cell(value)
    match = LEVEL_LABEL_RE.match(text)
    if not match:
        return "", text
    label = match.group(1).lower().replace(" ", "")
    if label == "program":
        return LEVEL_PROGRAM, text[match.end() :].strip(" :-\t")
    if label == "kegiatan":
        return LEVEL_KEGIATAN, text[match.end() :].strip(" :-\t")
    return LEVEL_SUBKEGIATAN, text[match.end() :].strip(" :-\t")


def clean_name(value: Any) -> str:
    """Bersihkan nama dari kode, nomor urut, dan label level."""
    text = clean_cell(value)
    if not text:
        return ""
    text = text.replace("\n", " ")
    code_match = CODE_RE.match(text)
    if code_match:
        text = text[code_match.end() :].strip(" :-\t")
    text = LEADING_NUMBER_RE.sub("", text).strip()
    _, text = strip_level_label(text)
    text = text.strip(" :-\t")
    return norm_space(text)


def level_from_code(code: str) -> str | None:
    parts = [part for part in code.split(".") if part]
    if len(parts) == 2:
        return LEVEL_PROGRAM
    if len(parts) == 3:
        return LEVEL_KEGIATAN
    if len(parts) >= 4:
        return LEVEL_SUBKEGIATAN
    return None


def parse_pagu(value: Any) -> float | None:
    """Parse nilai pagu Indonesia/Amerika menjadi float."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_cell(value)
    if not text:
        return None
    negative = text.strip().startswith("(") and text.strip().endswith(")")
    cleaned = PAGU_RE.sub("", text).replace("-", "")
    if not cleaned or cleaned in {".", ","}:
        return None
    if "." in cleaned and "," in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        groups = cleaned.split(",")
        if len(groups) > 2 or (len(groups[-1]) == 3 and len(groups[0]) <= 3):
            cleaned = cleaned.replace(",", "")
        else:
            cleaned = cleaned.replace(",", ".")
    elif "." in cleaned:
        groups = cleaned.split(".")
        if len(groups) > 2 or (len(groups[-1]) == 3 and len(groups[0]) <= 3 and len(groups) == 2):
            if len(groups[-1]) == 3 and len(groups[0]) <= 3:
                cleaned = cleaned.replace(".", "")
            else:
                cleaned = cleaned.replace(".", "")
        else:
            cleaned = cleaned
    try:
        number = float(cleaned)
    except ValueError:
        return None
    return -number if negative else number


def split_indicators(indicator_value: Any, target_value: Any, satuan_value: Any) -> list[dict[str, str]]:
    indicator = clean_cell(indicator_value)
    target = clean_cell(target_value)
    satuan = clean_cell(satuan_value)
    if not indicator and not target and not satuan:
        return []
    parts = [part.strip() for part in re.split(r"[\n;]+", indicator) if part.strip()]
    if not parts:
        return [{"nama": "", "target": target, "satuan": satuan}]
    if len(parts) > 1:
        targets = [part.strip() for part in re.split(r"[\n;]+", target) if part.strip()]
        satuans = [part.strip() for part in re.split(r"[\n;]+", satuan) if part.strip()]
        result = []
        for index, part in enumerate(parts):
            result.append(
                {
                    "nama": part,
                    "target": targets[index] if index < len(targets) else (targets[0] if len(targets) == 1 else ""),
                    "satuan": satuans[index] if index < len(satuans) else (satuans[0] if len(satuans) == 1 else ""),
                }
            )
        return result
    return [{"nama": parts[0], "target": target, "satuan": satuan}]


# ----------------------------------------------------------------------
# Ekstraksi file mentah
# ----------------------------------------------------------------------
def extract_any(path: str | Path, filename: str | None = None) -> dict[str, Any]:
    """Ambil teks dan tabel dari sebuah file."""
    target = Path(path)
    name = filename or target.name
    suffix = Path(name).suffix.lower() or target.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(target)
    if suffix == ".docx":
        return _extract_docx(target)
    if suffix == ".xlsx":
        return _extract_xlsx(target)
    if suffix == ".csv":
        return _extract_csv(target)
    if suffix in {".txt", ".md"}:
        text = _decode(target.read_bytes())
        return {"format": suffix.lstrip("."), "text": text, "tables": [], "warnings": []}
    raise ValueError(f"Format file tidak didukung: {suffix or 'tanpa ekstensi'}")


def _extract_pdf(path: Path) -> dict[str, Any]:
    import fitz  # PyMuPDF

    pages_text: list[str] = []
    tables: list[dict[str, Any]] = []
    warnings: list[str] = []
    with fitz.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            pages_text.append(page.get_text("text"))
            try:
                found = page.find_tables()
                for table_index, table in enumerate(found.tables, start=1):
                    rows = table.extract()
                    if rows:
                        tables.append({"sheet": f"Halaman {page_number} tabel {table_index}", "rows": rows})
            except Exception:
                # Versi PyMuPDF lama tidak memiliki find_tables; teks tetap dipakai.
                pass
    text = "\n\n".join(pages_text)
    if not text.strip():
        warnings.append(
            "PDF tidak mengandung teks (kemungkinan hasil pindai). Perlu OCR atau fallback LLM dengan teks hasil OCR."
        )
    else:
        warnings.append("Ekstraksi PDF berbasis teks/tabel; periksa kembali hasil pemetaan baris.")
    return {"format": "pdf", "text": text, "tables": tables, "warnings": warnings}


def _extract_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables: list[dict[str, Any]] = []
    for index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        if rows:
            tables.append({"sheet": f"Tabel {index}", "rows": rows})
    table_text = "\n".join(" | ".join(clean_cell(cell) for cell in row) for table in tables for row in table["rows"])
    text = "\n".join(paragraphs)
    if table_text:
        text = f"{text}\n\n{table_text}" if text else table_text
    return {"format": "docx", "text": text, "tables": tables, "warnings": []}


def _extract_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=False, data_only=True)
    tables: list[dict[str, Any]] = []
    text_parts: list[str] = []
    warnings: list[str] = []
    for worksheet in workbook.worksheets:
        max_row = min(worksheet.max_row or 0, 8000)
        max_col = min(worksheet.max_column or 0, 80)
        if max_row <= 0 or max_col <= 0:
            continue
        values = [[worksheet.cell(row=row, column=col).value for col in range(1, max_col + 1)] for row in range(1, max_row + 1)]
        for merged in list(worksheet.merged_cells.ranges):
            try:
                top_value = worksheet.cell(row=merged.min_row, column=merged.min_col).value
            except (ValueError, IndexError):
                continue
            for row in range(merged.min_row - 1, min(merged.max_row, max_row)):
                for col in range(merged.min_col - 1, min(merged.max_col, max_col)):
                    if values[row][col] in (None, ""):
                        values[row][col] = top_value
        if any(any(cell not in (None, "") for cell in row) for row in values):
            tables.append({"sheet": worksheet.title, "rows": values})
        for row in values[:2000]:
            line = "\t".join(clean_cell(cell) for cell in row if cell not in (None, ""))
            if line:
                text_parts.append(line)
        if max_row >= 8000:
            warnings.append(f"Sheet {worksheet.title} dipotong pada 8.000 baris.")
    workbook.close()
    return {"format": "xlsx", "text": "\n".join(text_parts), "tables": tables, "warnings": warnings}


def _extract_csv(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    text = _decode(raw)
    lines = text.splitlines()
    if lines and lines[0].lower().startswith("sep="):
        lines = lines[1:]
    delimiter = ","
    try:
        delimiter = csv.Sniffer().sniff("\n".join(lines[:20]), delimiters=",;|\t").delimiter
    except csv.Error:
        pass
    reader = csv.reader(io.StringIO("\n".join(lines)), delimiter=delimiter)
    rows = [list(row) for row in reader]
    text = "\n".join("\t".join(row) for row in rows)
    return {"format": "csv", "text": text, "tables": [{"sheet": "CSV", "rows": rows}], "warnings": []}


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# ----------------------------------------------------------------------
# Pemetaan tabel menjadi record
# ----------------------------------------------------------------------
def _match_header(cell: Any) -> str | None:
    text = clean_cell(cell)
    if not text:
        return None
    compact = text.lower().strip()
    for canonical in ("kode", "indikator", "target", "satuan", "pagu", "lokasi", "keterangan", "sasaran", "no"):
        if HEADER_PATTERNS[canonical].search(compact):
            return canonical
    if HEADER_PATTERNS["level"].search(compact):
        return "level"
    if HEADER_PATTERNS["nama"].search(compact):
        return "nama"
    if re.search(r"\b(program|kegiatan|subkegiatan|sub kegiatan)\b", compact):
        return "nama"
    return None


def _header_mapping(rows: list[list[Any]], start: int = 0) -> tuple[dict[int, str], int]:
    best_mapping: dict[int, str] = {}
    best_score = 0
    best_row = -1
    limit = min(len(rows), start + 30)
    for row_index in range(start, limit):
        row = rows[row_index]
        mapping: dict[int, str] = {}
        used: set[str] = set()
        for column_index, cell in enumerate(row):
            canonical = _match_header(cell)
            if canonical and canonical not in used:
                mapping[column_index] = canonical
                used.add(canonical)
        score = len(mapping)
        if "nama" in used or "kode" in used:
            score += 1
        if score > best_score:
            best_score = score
            best_mapping = mapping
            best_row = row_index
    if best_score < 3 or "nama" not in set(best_mapping.values()):
        # Coba gabungkan dua baris header (header bertingkat).
        for row_index in range(start, min(len(rows) - 1, start + 20)):
            combined = [
                f"{clean_cell(rows[row_index][i])} {clean_cell(rows[row_index + 1][i])}".strip()
                for i in range(max(len(rows[row_index]), len(rows[row_index + 1])))
            ]
            mapping = {}
            used = set()
            for column_index, cell in enumerate(combined):
                canonical = _match_header(cell)
                if canonical and canonical not in used:
                    mapping[column_index] = canonical
                    used.add(canonical)
            score = len(mapping) + (1 if "nama" in used else 0)
            if score > best_score and "nama" in used:
                best_score = score
                best_mapping = mapping
                best_row = row_index + 1
    if not best_mapping:
        return {}, -1
    # Fallback kolom nama: kolom teks terpanjang yang belum terpetakan.
    if "nama" not in set(best_mapping.values()):
        candidates: list[tuple[int, int]] = []
        for column_index in range(max(len(row) for row in rows[best_row + 1 : best_row + 15] or [[]])):
            if column_index in best_mapping:
                continue
            lengths = []
            for row in rows[best_row + 1 : best_row + 20]:
                if column_index < len(row):
                    text = clean_cell(row[column_index])
                    if text and not text.replace(".", "").isdigit():
                        lengths.append(len(text))
            if lengths:
                candidates.append((sum(lengths) // len(lengths), column_index))
        if candidates:
            candidates.sort(reverse=True)
            best_mapping[candidates[0][1]] = "nama"
    return best_mapping, best_row


def _rows_to_records(rows: list[list[Any]], source: str, warnings: list[str]) -> list[dict[str, Any]]:
    mapping, header_row = _header_mapping(rows)
    if header_row < 0 or "nama" not in set(mapping.values()):
        if any(any(clean_cell(cell) for cell in row) for row in rows):
            warnings.append(f"Tabel {source} tidak memiliki header kolom yang dikenali; dilewati.")
        return []
    records: list[dict[str, Any]] = []
    blanks = 0
    for row_index in range(header_row + 1, len(rows)):
        row = rows[row_index]
        values = {canonical: row[column] if column < len(row) else "" for column, canonical in mapping.items()}
        if not any(clean_cell(value) for value in values.values()):
            blanks += 1
            if blanks >= 5:
                break
            continue
        blanks = 0
        nama_raw = clean_cell(values.get("nama"))
        kode = extract_code(values.get("kode")) or extract_code(nama_raw)
        level = ""
        level_cell = clean_cell(values.get("level"))
        if level_cell:
            mapped_level, _ = strip_level_label(level_cell)
            level = mapped_level
            if not level:
                upper = level_cell.upper()
                if "SUB" in upper:
                    level = LEVEL_SUBKEGIATAN
                elif "KEGIATAN" in upper:
                    level = LEVEL_KEGIATAN
                elif "PROGRAM" in upper:
                    level = LEVEL_PROGRAM
        label_level, _ = strip_level_label(nama_raw)
        if not level:
            level = label_level
        if not level:
            level = level_from_code(kode) or ""
        nama = clean_name(nama_raw)
        if not nama and not kode and not values.get("indikator"):
            continue
        if SUMMARY_RE.match(nama) or SUMMARY_RE.match(nama_raw):
            continue
        if HEADERISH_RE.search(nama_raw) and len(HEADERISH_RE.findall(nama_raw)) >= 2:
            continue
        record = {
            "level": level,
            "kode": kode,
            "nama": nama,
            "indikator": split_indicators(values.get("indikator"), values.get("target"), values.get("satuan")),
            "pagu": parse_pagu(values.get("pagu")),
            "pagu_raw": clean_cell(values.get("pagu")),
            "lokasi": clean_cell(values.get("lokasi")),
            "keterangan": clean_cell(values.get("keterangan")),
            "sasaran": clean_cell(values.get("sasaran")),
            "source": source,
            "row": row_index + 1,
        }
        records.append(record)
    return records


def _table_to_records(table: dict[str, Any], warnings: list[str]) -> list[dict[str, Any]]:
    rows = table.get("rows") or []
    source = table.get("sheet") or "tabel"
    return _rows_to_records(rows, source, warnings)


def _records_to_programs(records: Iterable[dict[str, Any]], warnings: list[str]) -> list[dict[str, Any]]:
    programs: list[dict[str, Any]] = []
    current_program: dict[str, Any] | None = None
    current_kegiatan: dict[str, Any] | None = None
    current_sub: dict[str, Any] | None = None

    def new_program(kode: str, nama: str, record: dict[str, Any] | None = None) -> dict[str, Any]:
        program = {
            "kode": kode,
            "nama": nama,
            "indikator": [],
            "sasaran": [],
            "pagu": None,
            "kegiatan": [],
            "_sintetis": False,
            "sumber": (record or {}).get("source", ""),
        }
        programs.append(program)
        return program

    def ensure_program(kode: str, nama: str, record: dict[str, Any] | None = None) -> dict[str, Any]:
        nonlocal current_program
        if current_program and kode and current_program.get("kode") == kode:
            return current_program
        for program in programs:
            if kode and program.get("kode") == kode:
                current_program = program
                return program
        current_program = new_program(kode, nama, record)
        return current_program

    def new_kegiatan(program: dict[str, Any], kode: str, nama: str, record: dict[str, Any] | None = None) -> dict[str, Any]:
        kegiatan = {
            "kode": kode,
            "nama": nama,
            "indikator": [],
            "pagu": None,
            "lokasi": "",
            "keterangan": "",
            "subkegiatan": [],
            "_sintetis": False,
            "sumber": (record or {}).get("source", ""),
        }
        program["kegiatan"].append(kegiatan)
        return kegiatan

    def ensure_kegiatan(program: dict[str, Any], kode: str, nama: str, record: dict[str, Any] | None = None) -> dict[str, Any]:
        nonlocal current_kegiatan
        if current_kegiatan and kode and current_kegiatan.get("kode") == kode:
            return current_kegiatan
        for kegiatan in program.get("kegiatan", []):
            if kode and kegiatan.get("kode") == kode:
                current_kegiatan = kegiatan
                return kegiatan
        current_kegiatan = new_kegiatan(program, kode, nama, record)
        return current_kegiatan

    def new_sub(kegiatan: dict[str, Any], kode: str, nama: str, record: dict[str, Any] | None = None) -> dict[str, Any]:
        sub = {
            "kode": kode,
            "nama": nama,
            "indikator": [],
            "target": "",
            "satuan": "",
            "pagu": None,
            "lokasi": "",
            "keterangan": "",
            "_sintetis": False,
            "sumber": (record or {}).get("source", ""),
        }
        kegiatan["subkegiatan"].append(sub)
        return sub

    def ensure_sub(kegiatan: dict[str, Any], kode: str, nama: str, record: dict[str, Any] | None = None) -> dict[str, Any]:
        nonlocal current_sub
        if current_sub and kode and current_sub.get("kode") == kode:
            return current_sub
        for sub in kegiatan.get("subkegiatan", []):
            if kode and sub.get("kode") == kode:
                current_sub = sub
                return sub
        current_sub = new_sub(kegiatan, kode, nama, record)
        return current_sub

    def merge_indicator(item: dict[str, Any], indicators: list[dict[str, str]]) -> None:
        for indicator in indicators:
            if not indicator.get("nama") and not indicator.get("target") and not indicator.get("satuan"):
                continue
            existing = None
            for candidate in item.get("indikator", []):
                if candidate.get("nama", "").lower() == indicator.get("nama", "").lower():
                    existing = candidate
                    break
            if existing:
                for key in ("target", "satuan"):
                    if not existing.get(key) and indicator.get(key):
                        existing[key] = indicator[key]
            else:
                item.setdefault("indikator", []).append(dict(indicator))

    def apply_meta(item: dict[str, Any], record: dict[str, Any]) -> None:
        if record.get("pagu") is not None and item.get("pagu") is None:
            item["pagu"] = record["pagu"]
        if record.get("lokasi") and not item.get("lokasi"):
            item["lokasi"] = record["lokasi"]
        if record.get("keterangan") and not item.get("keterangan"):
            item["keterangan"] = record["keterangan"]
        if record.get("sasaran") and item.get("sasaran") is not None and not item.get("sasaran"):
            item["sasaran"] = [record["sasaran"]]
        merge_indicator(item, record.get("indikator", []))

    for record in records:
        level = record.get("level") or ""
        kode = record.get("kode") or ""
        nama = record.get("nama") or ""
        indicators = record.get("indikator") or []

        if not level:
            if indicators:
                target_item = current_sub or current_kegiatan or current_program
                if target_item is not None:
                    apply_meta(target_item, record)
                    continue
            if nama and current_sub and not current_sub.get("nama"):
                current_sub["nama"] = nama
                continue
            if nama and current_kegiatan and not current_kegiatan.get("nama"):
                current_kegiatan["nama"] = nama
                continue
            if nama and current_program and not current_program.get("nama"):
                current_program["nama"] = nama
                continue
            if kode and not nama:
                continue
            if nama:
                warnings.append(f"Baris tanpa level/kode dilewati: \"{nama[:80]}\".")
            continue

        if level == LEVEL_PROGRAM:
            if current_program and not kode and nama and current_program.get("kode"):
                # Nama program pada baris lanjutan.
                if not current_program.get("nama"):
                    current_program["nama"] = nama
                    apply_meta(current_program, record)
                continue
            program = ensure_program(kode, nama, record)
            apply_meta(program, record)
            current_program = program
            current_kegiatan = None
            current_sub = None
            continue

        if level == LEVEL_KEGIATAN:
            program = current_program
            if program is None:
                parent_code = ".".join(kode.split(".")[:2]) if kode else ""
                program = ensure_program(parent_code, "", record)
                program["_sintetis"] = True
                warnings.append(f"Program induk untuk kegiatan {kode or nama} tidak ditemukan; dibuat penampung sementara.")
            kegiatan = ensure_kegiatan(program, kode, nama, record)
            apply_meta(kegiatan, record)
            current_kegiatan = kegiatan
            current_sub = None
            continue

        if level == LEVEL_SUBKEGIATAN:
            if current_program is None:
                parent_program_code = ".".join(kode.split(".")[:2]) if kode else ""
                current_program = ensure_program(parent_program_code, "", record)
                current_program["_sintetis"] = True
                warnings.append(f"Program induk untuk subkegiatan {kode or nama} tidak ditemukan; dibuat penampung sementara.")
            if current_kegiatan is None:
                parent_kegiatan_code = ".".join(kode.split(".")[:3]) if kode else ""
                current_kegiatan = ensure_kegiatan(current_program, parent_kegiatan_code, "", record)
                current_kegiatan["_sintetis"] = True
                warnings.append(f"Kegiatan induk untuk subkegiatan {kode or nama} tidak ditemukan; dibuat penampung sementara.")
            sub = ensure_sub(current_kegiatan, kode, nama, record)
            apply_meta(sub, record)
            if sub.get("indikator"):
                first = sub["indikator"][0]
                if not sub.get("target"):
                    sub["target"] = first.get("target", "")
                if not sub.get("satuan"):
                    sub["satuan"] = first.get("satuan", "")
            current_sub = sub
            continue

    return programs


# ----------------------------------------------------------------------
# Parsing teks (fallback PDF/teks)
# ----------------------------------------------------------------------
def _text_to_records(text: str, warnings: list[str]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    last_record: dict[str, Any] | None = None
    lines = [norm_space(line) for line in text.splitlines()]
    header_map: dict[int, str] | None = None
    for line in lines:
        if not line:
            continue
        if "|" in line or "\t" in line:
            cells = [cell.strip() for cell in re.split(r"[|\t]", line) if cell.strip()]
            if header_map is None and sum(1 for cell in cells if _match_header(cell)) >= 2:
                header_map = {index: mapped for index, cell in enumerate(cells) if (mapped := _match_header(cell))}
                continue
            if header_map and cells:
                values = {
                    canonical: cells[index] if index < len(cells) else ""
                    for index, canonical in header_map.items()
                }
                kode = extract_code(values.get("kode")) or extract_code(values.get("nama"))
                level = level_from_code(kode) or ""
                label_level, _ = strip_level_label(values.get("nama"))
                level = label_level or level
                nama = clean_name(values.get("nama"))
                if nama and not SUMMARY_RE.match(nama):
                    record = {
                        "level": level,
                        "kode": kode,
                        "nama": nama,
                        "indikator": split_indicators(values.get("indikator"), values.get("target"), values.get("satuan")),
                        "pagu": parse_pagu(values.get("pagu")),
                        "pagu_raw": clean_cell(values.get("pagu")),
                        "lokasi": clean_cell(values.get("lokasi")),
                        "keterangan": clean_cell(values.get("keterangan")),
                        "sasaran": clean_cell(values.get("sasaran")),
                        "source": "teks",
                        "row": 0,
                    }
                    records.append(record)
                    last_record = record
                continue
        code_match = re.search(r"(?<!\d)(\d{1,2}(?:\.\d{1,2}){1,4})(?!\d)", line)
        if code_match:
            before = line[: code_match.start()].strip()
            after = line[code_match.end() :].strip(" .:-\t")
            if not after:
                continue
            if len(before) > 5 and not re.fullmatch(r"(kode|kd|no|nomor|[\d.]+)", before, re.I):
                continue
            kode = code_match.group(1)
            level = level_from_code(kode) or ""
            label_level, after = strip_level_label(after)
            level = label_level or level
            nama = clean_name(after)
            if not nama or SUMMARY_RE.match(nama):
                continue
            record = {
                "level": level,
                "kode": kode,
                "nama": nama,
                "indikator": [],
                "pagu": None,
                "pagu_raw": "",
                "lokasi": "",
                "keterangan": "",
                "sasaran": "",
                "source": "teks",
                "row": 0,
            }
            records.append(record)
            last_record = record
            continue
        label_level, rest = strip_level_label(line)
        if label_level and rest:
            kode = extract_code(rest)
            nama = clean_name(rest)
            if nama:
                record = {
                    "level": label_level,
                    "kode": kode,
                    "nama": nama,
                    "indikator": [],
                    "pagu": None,
                    "pagu_raw": "",
                    "lokasi": "",
                    "keterangan": "",
                    "sasaran": "",
                    "source": "teks",
                    "row": 0,
                }
                records.append(record)
                last_record = record
                continue
        indicator_match = re.match(r"^indikator\s*[:\-]\s*(.+)", line, re.I)
        target_match = re.match(r"^(?:target|satuan)\s*[:\-]\s*(.+)", line, re.I)
        if indicator_match and last_record is not None:
            last_record["indikator"].append({"nama": indicator_match.group(1), "target": "", "satuan": ""})
        elif target_match and last_record is not None and last_record["indikator"]:
            key = "target" if line.lower().startswith("target") else "satuan"
            if not last_record["indikator"][-1].get(key):
                last_record["indikator"][-1][key] = target_match.group(1).strip()
    if not records:
        warnings.append(
            "Tidak ada baris Program/Kegiatan/Subkegiatan yang dapat diparse dari teks. "
            "Gunakan fallback LLM atau unggah dokumen XLSX/DOCX yang tabelnya terbaca."
        )
    return records


# ----------------------------------------------------------------------
# Identitas dokumen
# ----------------------------------------------------------------------
def detect_identity(text: str, filename: str) -> dict[str, Any]:
    compact = norm_space(text)
    upper = compact.upper()
    opd = ""
    kode_opd = ""
    opd_patterns = [
        re.compile(r"(?:NAMA\s+)?(?:ORGANISASI\s+PERANGKAT\s+DAERAH|OPD|PERANGKAT\s+DAERAH|SKPD)\s*[:\-]\s*([^\n]{3,140})", re.I),
        re.compile(r"(?:KODE\s+OPD|KODE\s+PERANGKAT\s+DAERAH)\s*[:\-]\s*([\d.]+)", re.I),
    ]
    for pattern in opd_patterns:
        match = pattern.search(compact)
        if match:
            value = norm_space(match.group(1))
            if value and len(value) < 160:
                if "KODE" in pattern.pattern:
                    kode_opd = value
                else:
                    opd = value
                    break
    if not opd:
        known = [
            "BADAN PERENCANAAN PEMBANGUNAN DAERAH",
            "BADAN PERENCANAAN PEMBANGUNAN DAERAH KABUPATEN SOPPENG",
            "BAPPEDA",
        ]
        for candidate in known:
            if candidate in upper:
                opd = candidate
                break
    if not opd and "bappeda" in filename.lower():
        opd = "BADAN PERENCANAAN PEMBANGUNAN DAERAH KABUPATEN SOPPENG"
    if not opd:
        for line in text.splitlines()[:80]:
            line_s = norm_space(line)
            if not line_s or "\t" in line or "|" in line or len(line_s) > 160:
                continue
            match = re.match(
                r"^(BADAN|DINAS|SEKRETARIAT|KANTOR|RUMAH SAKIT|KECAMATAN|INSPEKTORAT)\s+[A-Z][A-Z0-9 /\.\-]{4,120}$",
                line_s.upper(),
            )
            if match:
                opd = line_s
                break
    tahun = ""
    for pattern in (r"TAHUN\s+ANGGARAN\s+(20\d{2})", r"TAHUN\s+(20\d{2})", r"\b(20\d{2})\b"):
        match = re.search(pattern, upper)
        if match:
            tahun = match.group(1)
            break
    judul = ""
    for line in text.splitlines()[:80]:
        line_s = norm_space(line)
        if "\t" in line or "|" in line or len(line_s) > 200:
            continue
        if re.search(r"RENCANA\s+KERJA|RENJA|RENCANA\s+KINERJA", line_s, re.I):
            judul = line_s
            break
    if not judul:
        judul = f"Rencana Kerja (Renja) {opd or 'OPD'}" + (f" Tahun {tahun}" if tahun else "")
    opd = re.sub(r"\s*[:\-]\s*$", "", opd).strip()
    if opd.upper() in {"BAPPEDA", "BADAN PERENCANAAN PEMBANGUNAN DAERAH"}:
        opd = "Badan Perencanaan Pembangunan Daerah Kabupaten Soppeng"
    return {
        "opd": opd,
        "kode_opd": kode_opd,
        "tahun": tahun,
        "judul": judul,
        "nama_file": Path(filename).name,
    }


# ----------------------------------------------------------------------
# Entry point
# ----------------------------------------------------------------------
def extract_renja(
    path: str | Path,
    filename: str | None = None,
    use_llm: bool | None = None,
    master: Any = None,
) -> dict[str, Any]:
    """Ekstrak satu dokumen Renja menjadi struktur terstruktur."""
    target = Path(path)
    name = filename or target.name
    raw = extract_any(target, name)
    warnings = list(raw.get("warnings", []))
    text = raw.get("text", "")

    records: list[dict[str, Any]] = []
    for table in raw.get("tables", []):
        table_records = _table_to_records(table, warnings)
        records.extend(table_records)

    # Deduplikasi record identik yang muncul karena header berulang.
    deduped: list[dict[str, Any]] = []
    seen: set[tuple] = set()
    for record in records:
        signature = (
            record.get("level"),
            record.get("kode"),
            record.get("nama"),
            tuple((item.get("nama"), item.get("target"), item.get("satuan")) for item in record.get("indikator", [])),
            record.get("pagu"),
        )
        if signature in seen:
            continue
        seen.add(signature)
        deduped.append(record)
    records = deduped

    programs = _records_to_programs(records, warnings)
    extraction_source = "tabel" if programs else ""

    if not programs and text:
        text_records = _text_to_records(text, warnings)
        if text_records:
            programs = _records_to_programs(text_records, warnings)
            extraction_source = "teks"

    identity = detect_identity(text, name)
    doc = {
        "opd": {
            "nama": identity["opd"],
            "kode": identity["kode_opd"],
        },
        "tahun": identity["tahun"],
        "dokumen": {
            "jenis": "Renja / Rencana Kinerja OPD",
            "judul": identity["judul"],
            "file": Path(name).name,
            "format": raw.get("format", Path(name).suffix.lstrip(".")),
            "sumber_ekstraksi": extraction_source or "tidak_terbaca",
            "diekstrak_pada": utc_now_iso(),
        },
        "program": programs,
        "warnings": _unique(warnings),
        "meta": {
            "jumlah_record": len(records),
            "sumber_tabel": len(raw.get("tables", [])),
            "panjang_teks": len(text),
        },
    }

    if use_llm is None:
        use_llm = bool(os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY"))
    if use_llm and master is not None and not programs:
        try:
            import llm_analyzer

            if llm_analyzer.is_enabled():
                llm_doc = llm_analyzer.extract_document(text, master, filename=name)
                if llm_doc.get("program"):
                    doc["program"] = llm_doc["program"]
                    doc["dokumen"]["sumber_ekstraksi"] = "llm"
                    doc["warnings"].append("Struktur Renja diambil melalui fallback LLM; wajib diverifikasi manual.")
                    if llm_doc.get("opd") and not doc["opd"]["nama"]:
                        doc["opd"]["nama"] = llm_doc["opd"]
                    if llm_doc.get("tahun") and not doc["tahun"]:
                        doc["tahun"] = llm_doc["tahun"]
        except Exception as exc:  # pragma: no cover - bergantung API eksternal
            doc["warnings"].append(f"Fallback LLM gagal: {exc}")
    elif use_llm and not programs:
        doc["warnings"].append("Fallback LLM tidak dijalankan karena master/dokumen belum siap.")

    doc["statistik"] = document_statistics(doc)
    doc["warnings"] = _unique(doc["warnings"])
    return doc


def extract_reference(path: str | Path, filename: str | None = None) -> dict[str, Any]:
    """Ekstrak teks dokumen acuan (RPJMD/Renstra/RKPD) untuk penyelarasan."""
    target = Path(path)
    name = filename or target.name
    raw = extract_any(target, name)
    text = raw.get("text", "")
    if len(text) > 800_000:
        text = text[:800_000]
    return {
        "nama_file": Path(name).name,
        "format": raw.get("format", ""),
        "teks": text,
        "jumlah_karakter": len(text),
        "warnings": raw.get("warnings", []),
        "diekstrak_pada": utc_now_iso(),
    }


def document_statistics(doc: dict[str, Any]) -> dict[str, int]:
    programs = doc.get("program") or []
    kegiatan = sum(len(program.get("kegiatan", [])) for program in programs)
    subkegiatan = sum(len(kegiatan.get("subkegiatan", [])) for program in programs for kegiatan in program.get("kegiatan", []))
    indikator = 0
    pagu = 0.0
    for program in programs:
        indikator += len(program.get("indikator", []))
        for kegiatan_item in program.get("kegiatan", []):
            indikator += len(kegiatan_item.get("indikator", []))
            for sub in kegiatan_item.get("subkegiatan", []):
                indikator += len(sub.get("indikator", []))
                if isinstance(sub.get("pagu"), (int, float)):
                    pagu += float(sub["pagu"])
    return {
        "program": len(programs),
        "kegiatan": kegiatan,
        "subkegiatan": subkegiatan,
        "indikator": indikator,
        "pagu": round(pagu, 2),
    }


def _unique(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = norm_space(value)
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return result


def iter_subkegiatan(doc: dict[str, Any]):
    for program in doc.get("program", []):
        for kegiatan in program.get("kegiatan", []):
            for sub in kegiatan.get("subkegiatan", []):
                yield program, kegiatan, sub


def iter_all_items(doc: dict[str, Any]):
    for program in doc.get("program", []):
        yield LEVEL_PROGRAM, program, None, None
        for kegiatan in program.get("kegiatan", []):
            yield LEVEL_KEGIATAN, kegiatan, program, None
            for sub in kegiatan.get("subkegiatan", []):
                yield LEVEL_SUBKEGIATAN, sub, program, kegiatan
