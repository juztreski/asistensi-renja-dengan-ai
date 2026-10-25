"""Membuat contoh Renja Bappeda dalam format XLSX untuk uji coba aplikasi."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT = BASE_DIR / "sample" / "renja_bappeda_contoh.xlsx"

NAVY = "16324F"
LIGHT_BLUE = "DCE6F1"
LIGHT_TEAL = "D9EFEC"
WHITE = "FFFFFF"
GRAY = "667085"
BORDER_COLOR = "C9D2DC"

HEADERS = [
    "No",
    "Level",
    "Kode",
    "Uraian / Nama Program, Kegiatan, dan Subkegiatan",
    "Indikator Kinerja",
    "Target",
    "Satuan",
    "Pagu Indikatif (Rp)",
    "Lokasi",
    "Keterangan",
]

ROWS = [
    ("Program", "1.01", "Program Penunjang Urusan Pemerintahan Daerah Kabupaten/Kota",
     "Persentase layanan administrasi perkantoran yang terpenuhi", "100", "Persen", None,
     "Kabupaten Soppeng", ""),
    ("Kegiatan", "1.01.01", "Perencanaan, Penganggaran, dan Evaluasi Kinerja Perangkat Daerah",
     "Jumlah dokumen perencanaan dan laporan kinerja yang tersusun", "5", "Dokumen", None,
     "Kabupaten Soppeng", ""),
    ("Subkegiatan", "1.01.01.2.01", "Penyusunan Dokumen Perencanaan Perangkat Daerah",
     "Jumlah dokumen perencanaan perangkat daerah", "1", "Dokumen", 15000000,
     "Kabupaten Soppeng", "Dokumen Renja dan RKA"),
    ("Subkegiatan", "1.01.01.2.02", "Penyusunan Anggaran Perangkat Daerah",
     "Jumlah dokumen anggaran perangkat daerah", "1", "Dokumen", 10000000,
     "Kabupaten Soppeng", ""),
    ("Subkegiatan", "1.01.01.2.03", "Evaluasi Kinerja Perangkat Daerah",
     "Jumlah laporan evaluasi kinerja perangkat daerah", "4", "Laporan", None,
     "Kabupaten Soppeng", "Pagu belum diisi"),
    ("Kegiatan", "1.01.03", "Administrasi Umum Perangkat Daerah",
     "Persentase kebutuhan administrasi umum yang terpenuhi", "100", "Persen", None,
     "Kabupaten Soppeng", ""),
    ("Subkegiatan", "1.01.03.2.02", "Penyediaan Alat Tulis Kantor",
     "Jumlah paket alat tulis kantor yang tersedia", "1", "Paket", 5000000,
     "Kabupaten Soppeng", "Nama tidak sesuai master"),
    ("Subkegiatan", "1.01.03.2.04", "Penyediaan Bahan Logistik Kantor",
     "Jumlah bulan penyediaan bahan logistik", "12", "Bulan", 12000000,
     "Kabupaten Soppeng", ""),
    ("Subkegiatan", "1.01.03.2.09", "Penyediaan Jasa Kebersihan Kantor",
     "Jumlah bulan layanan kebersihan kantor", "12", "Bulan", 8000000,
     "Kabupaten Soppeng", "Kode tidak ada di master"),
    ("Program", "6.01", "Program Perencanaan, Pengendalian dan Evaluasi Pembangunan Daerah",
     "Persentase dokumen perencanaan daerah yang tersusun tepat waktu", "100", "Persen", None,
     "Kabupaten Soppeng", ""),
    ("Kegiatan", "6.01.01", "Perencanaan Pengendalian dan Evaluasi Pembangunan Daerah",
     "Jumlah dokumen perencanaan dan evaluasi pembangunan daerah", "6", "Dokumen", None,
     "Kabupaten Soppeng", ""),
    ("Subkegiatan", "6.01.01.2.01", "Penyusunan Dokumen Perencanaan Pembangunan Daerah",
     "Jumlah dokumen perencanaan pembangunan daerah", "1", "Dokumen", 50000000,
     "Kabupaten Soppeng", ""),
    ("Subkegiatan", "6.01.01.2.02", "Koordinasi Perencanaan Pembangunan Daerah",
     "", "4", "Kali", 25000000, "Kabupaten Soppeng", "Indikator belum diisi"),
    ("Kegiatan", "6.01.02", "Koordinasi dan Sinkronisasi Perencanaan Pembangunan Daerah",
     "Jumlah kegiatan koordinasi perencanaan yang terlaksana", "12", "Kali", None,
     "Kabupaten Soppeng", ""),
    ("Subkegiatan", "6.01.02.2.01", "Koordinasi Penyusunan Dokumen Perencanaan Pembangunan Daerah",
     "Jumlah koordinasi penyusunan dokumen perencanaan", "5", "", 15000000,
     "Kabupaten Soppeng", "Satuan belum diisi"),
    ("Subkegiatan", "6.01.02.2.02", "Sinkronisasi Perencanaan Pembangunan Daerah",
     "Jumlah sinkronisasi perencanaan pembangunan", "TBD", "", 0,
     "Kabupaten Soppeng", "Target belum jelas dan pagu nol"),
]


def build() -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Renja Bappeda 2027"

    sheet.merge_cells("A1:J1")
    sheet["A1"] = "RENCANA KERJA PERANGKAT DAERAH (RENJA)"
    sheet["A1"].font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:J2")
    sheet["A2"] = "PEMERINTAH KABUPATEN SOPPENG - BADAN PERENCANAAN PEMBANGUNAN DAERAH"
    sheet["A2"].font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    sheet["A2"].fill = PatternFill("solid", fgColor="1F4E6B")
    sheet["A2"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[2].height = 21

    sheet.merge_cells("A3:J3")
    sheet["A3"] = "CONTOH DOKUMEN UJI COBA - memuat kesalahan yang disengaja untuk demonstrasi asistensi"
    sheet["A3"].font = Font(name="Calibri", size=9, italic=True, color=GRAY)
    sheet["A3"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[3].height = 18

    sheet["A5"] = "Nama OPD"
    sheet["B5"] = "Bappeda Kabupaten Soppeng"
    sheet["D5"] = "Tahun Anggaran"
    sheet["E5"] = 2027
    sheet["G5"] = "Tanggal"
    sheet["H5"] = "2026-10-24"
    for cell in ("A5", "D5", "G5"):
        sheet[cell].font = Font(name="Calibri", size=10, bold=True)
    for cell in ("B5", "E5", "H5"):
        sheet[cell].font = Font(name="Calibri", size=10)

    header_row = 7
    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(name="Calibri", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[header_row].height = 42

    for offset, row_data in enumerate(ROWS, start=1):
        row_number = header_row + offset
        level, code, name, indicator, target, unit, budget, location, note = row_data
        values = [offset, level, code, name, indicator, target, unit, budget, location, note]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.font = Font(name="Calibri", size=9, bold=(level == "Program"))
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                horizontal="center" if column in {1, 2, 6, 7} else "left",
                wrap_text=column in {4, 5, 8, 9, 10},
            )
            if level == "Program":
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            elif level == "Kegiatan":
                cell.fill = PatternFill("solid", fgColor=LIGHT_TEAL)
        if isinstance(budget, (int, float)):
            sheet.cell(row=row_number, column=8).number_format = "#,##0"
        sheet.row_dimensions[row_number].height = 30 if level == "Program" else 26

    total_rows = header_row + len(ROWS)
    sheet.merge_cells(start_row=total_rows + 2, start_column=1, end_row=total_rows + 2, end_column=10)
    note_cell = sheet.cell(
        row=total_rows + 2,
        column=1,
        value=(
            "Contoh ini digunakan untuk menguji ekstraksi tabel, pemeriksaan nomenklatur, kelengkapan indikator, "
            "pagu, dan penyelarasan. Bukan dokumen Renja resmi."
        ),
    )
    note_cell.font = Font(name="Calibri", size=9, italic=True, color=GRAY)
    note_cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[total_rows + 2].height = 26

    widths = {1: 5, 2: 14, 3: 18, 4: 46, 5: 42, 6: 10, 7: 12, 8: 18, 9: 20, 10: 30}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.freeze_panes = "A8"
    sheet.auto_filter.ref = f"A{header_row}:J{total_rows}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"1:{header_row}"
    sheet.oddFooter.center.text = "Halaman &P dari &N"
    sheet.oddFooter.center.size = 8

    workbook.properties.title = "Contoh Renja Bappeda Kabupaten Soppeng 2027"
    workbook.properties.subject = "Dokumen uji coba Asistensi Renja OPD"
    workbook.properties.creator = "Sistem Asistensi Bappeda Soppeng"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Contoh XLSX dibuat: {path} ({path.stat().st_size:,} byte)")
