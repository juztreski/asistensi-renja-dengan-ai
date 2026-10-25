from __future__ import annotations

import io

from docx import Document
from openpyxl import load_workbook

from align_check import analyze_document, recompute_after_review
from extract_renja import document_statistics, extract_reference, extract_renja
from generate_worksheet import build_docx, build_xlsx
from review_fields import apply_review, compute_score


def _reference(sample_dir, key: str, filename: str) -> dict:
    data = extract_reference(sample_dir / filename, filename)
    return {
        "jenis": key,
        "nama": key.upper(),
        "nama_file": filename,
        "teks": data["teks"],
        "jumlah_karakter": data["jumlah_karakter"],
        "warnings": data.get("warnings", []),
    }


def _prepare(sample_dir, master):
    doc = extract_renja(
        sample_dir / "renja_bappeda_contoh.xlsx",
        "renja_bappeda_contoh.xlsx",
        use_llm=False,
        master=master,
    )
    doc["tahun"] = "2027"
    doc["opd"] = {"nama": "Bappeda Kabupaten Soppeng"}
    doc["statistik"] = document_statistics(doc)
    references = [
        _reference(sample_dir, "rkpd", "rkpd_2027_cuplikan.txt"),
        _reference(sample_dir, "renstra", "renstra_bappeda_cuplikan.txt"),
    ]
    analysis = analyze_document(doc, master, references, extraction_warnings=doc.get("warnings", []))
    recomputed = recompute_after_review(analysis["findings"])
    analysis["score_final"] = recomputed["score"]
    analysis["summary_final"] = recomputed["summary"]
    return doc, references, analysis


def test_master_seed_counts_and_nomenclature(master):
    counts = master.counts()
    assert counts == {"program": 3, "kegiatan": 9, "subkegiatan": 36}
    assert master.is_sample is True

    exact = master.resolve("1.01.01.2.01", "Penyusunan Dokumen Perencanaan Perangkat Daerah", level="Subkegiatan")
    assert exact["status"] == "sesuai"

    wrong_name = master.resolve("1.01.03.2.02", "Penyediaan Alat Tulis Kantor", level="Subkegiatan")
    assert wrong_name["status"] != "sesuai"
    assert wrong_name["tingkat"] == "Kritis"


def test_extract_sample_xlsx(sample_dir, master):
    doc = extract_renja(
        sample_dir / "renja_bappeda_contoh.xlsx",
        "renja_bappeda_contoh.xlsx",
        use_llm=False,
        master=master,
    )
    assert document_statistics(doc) == {
        "program": 2,
        "kegiatan": 4,
        "subkegiatan": 10,
        "indikator": 16,
        "pagu": 140000000.0,
    }
    assert "Kabupaten Soppeng" in doc["opd"]["nama"]


def test_analysis_finds_intentional_errors(sample_dir, master):
    doc, _references, analysis = _prepare(sample_dir, master)
    findings = analysis["findings"]
    kinds = {finding["jenis"] for finding in findings}

    assert len(findings) == 17
    assert analysis["summary"]["by_severity"] == {"Kritis": 2, "Sedang": 6, "Saran": 9}
    assert 0 < analysis["score"]["score"] < 100
    assert {
        "NOMENKLATUR_TIDAK_SESUAI",
        "NOMENKLATUR_TIDAK_TERDAFTAR",
        "PAGU_KOSONG",
        "PAGU_NOL",
        "NAMA_INDIKATOR_KOSONG",
        "SATUAN_KOSONG",
        "TARGET_TIDAK_JELAS",
    } <= kinds
    assert len(analysis["alignment"]) == 2
    assert all(item["status"] == "Selaras Sebagian" for item in analysis["alignment"])


def test_review_can_resolve_all_findings(sample_dir, master):
    _doc, _references, analysis = _prepare(sample_dir, master)
    form = {f"status_{finding['id']}": "Dikoreksi" for finding in analysis["findings"]}
    updated, changed = apply_review(analysis["findings"], form)

    assert changed == len(analysis["findings"])
    recomputed = recompute_after_review(updated)
    assert recomputed["score"]["score"] == 100
    assert recomputed["summary"]["active"] == 0
    assert recomputed["summary"]["resolved"] == len(updated)
    assert compute_score(updated)["penalty"] == 0


def test_worksheet_generation_is_readable(sample_dir, master):
    doc, references, analysis = _prepare(sample_dir, master)
    job = {
        "doc": doc,
        "references": references,
        "analysis": analysis,
        "identity": {
            "opd": "Bappeda Kabupaten Soppeng",
            "tahun": "2027",
            "tanggal": "2026-10-24",
            "nomor": "LKA/001/X/2026",
            "perencana": "Perencana Ahli Pertama",
            "nip_perencana": "19800101 200501 1 001",
            "koordinator": "Koordinator Bidang",
            "nip_koordinator": "19750101 200001 1 001",
            "kepala": "Kepala Bappeda",
            "nip_kepala": "19700101 199501 1 001",
        },
    }

    docx_buffer = build_docx(job)
    docx_bytes = docx_buffer.getvalue()
    assert docx_bytes[:2] == b"PK"
    document = Document(io.BytesIO(docx_bytes))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "LEMBAR KERJA ASISTENSI" in text.upper()
    assert len(document.tables) >= 3

    xlsx_buffer = build_xlsx(job)
    xlsx_bytes = xlsx_buffer.getvalue()
    assert xlsx_bytes[:2] == b"PK"
    workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    assert {"Identitas", "Ringkasan", "Matriks", "Temuan", "Penyelarasan", "Paraf"} <= set(workbook.sheetnames)
    assert workbook["Ringkasan"]["C4"].value == analysis["score_final"]["score"]
    workbook.close()
